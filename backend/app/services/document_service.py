"""Сервис подписания документов простым подтверждением паролем."""

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import NotFoundError, UnauthorizedError
from app.core.security import verify_password
from app.models.main.document_file import DocumentFile
from app.models.main.user import User
from app.schemas.document import DocumentSignResponse


class DocumentService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def sign_document(
        self,
        *,
        current_user: User,
        document_id: int,
        password: str,
    ) -> DocumentSignResponse:
        if not verify_password(password, current_user.password_hash):
            raise UnauthorizedError("Неверный пароль для электронной подписи")

        result = await self.session.execute(
            select(DocumentFile)
            .options(selectinload(DocumentFile.request))
            .where(DocumentFile.id == document_id)
        )
        document_file = result.scalar_one_or_none()
        if document_file is None or document_file.request.employee_id != current_user.id:
            raise NotFoundError("Документ не найден")

        signed_at = datetime.now(timezone.utc)
        path = Path(document_file.file_path)
        if path.suffix.lower() == ".docx" and path.exists():
            document = Document(str(path))
            position = current_user.position.name if current_user.position else "Сотрудник"
            stamp = document.add_table(rows=1, cols=1)
            cell = stamp.cell(0, 0)
            cell.text = (
                "ДОКУМЕНТ ПОДПИСАН\n"
                "ПРОСТОЙ ЭЛЕКТРОННОЙ ПОДПИСЬЮ\n\n"
                f"Подписал: {current_user.full_name}\n"
                f"{position}\n"
                f"Дата: {signed_at.astimezone().strftime('%d.%m.%Y %H:%M')}"
            )
            document.save(str(path))

        return DocumentSignResponse(
            document_id=document_id,
            status="signed",
            message="Документ подписан простой электронной подписью",
            signed_at=signed_at,
        )
